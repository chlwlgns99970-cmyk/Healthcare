"""Explicit new endpoints only; retain original bytes and 90-dish page index."""
import hashlib,json,re,urllib.request,concurrent.futures
from urllib.parse import quote
from pathlib import Path
from lxml import html
from pypdf import PdfReader
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'data-source/recipe-residual-90'
def load(p):return json.loads(p.read_text(encoding='utf-8'))
def run():
 queue=load(OUT/'queue-90.json');targets=load(OUT/'source-targets.json');rawdir=OUT/'raw';rawdir.mkdir(exist_ok=True)
 old={r['url']:r for r in load(OUT/'source-captures.json')} if (OUT/'source-captures.json').exists() else {}
 def fetch(target):
  u=target['url'];r=dict(target,checkedAt='2026-10-06',status='FAILED',matches=[])
  try:
   previous=old.get(u);p=ROOT/previous['rawFile'] if previous and previous.get('rawFile') else None
   if p and p.exists() and hashlib.sha256(p.read_bytes()).hexdigest()==previous['sha256']:data=p.read_bytes();r['captureReused']=True
   else:
    with urllib.request.urlopen(urllib.request.Request(quote(u,safe=':/?=&%'),headers={'User-Agent':'Mozilla/5.0'}),timeout=30) as response:data=response.read();r['contentType']=response.headers.get('Content-Type','')
    p=rawdir/hashlib.sha256(u.encode()).hexdigest()[:20];p.write_bytes(data);r['captureReused']=False
   r.update(rawFile=p.relative_to(ROOT).as_posix(),sha256=hashlib.sha256(data).hexdigest(),bytes=len(data),status='CAPTURED')
   if len(data)<500 and (b'alert(' in data or b'location' in data):
    r['status']='ACCESS_REQUIRED_NO_DOCUMENT'
    return r
   if data.startswith(b'%PDF'):pages=[page.extract_text() or '' for page in PdfReader(p).pages]
   elif data.startswith(b'PK'):
    import io
    book=load_workbook(io.BytesIO(data),data_only=True,read_only=True)
    pages=[s.title+'\n'+'\n'.join(' | '.join(str(v) if v is not None else '' for v in row) for row in s.iter_rows(values_only=True)) for s in book]
    r['sheetNames']=book.sheetnames
    book.close()
   else:
    text=data.decode('utf-8',errors='replace');doc=html.fromstring(text)
    for node in doc.xpath('//script|//style'):node.drop_tree()
    pages=[doc.text_content()+' '+' '.join(doc.xpath('//img/@alt'))]
   textfile=OUT/(p.name+'.txt');textfile.write_text('\n\n'.join('PAGE '+str(n+1)+'\n'+s for n,s in enumerate(pages)),encoding='utf-8')
   r.update(textFile=textfile.relative_to(ROOT).as_posix(),pageCount=len(pages))
   compact=lambda s:re.sub(r'[^가-힣a-zA-Z0-9]','',s)
   cp=[compact(s) for s in pages]
   for food in queue:
    variants=[food['name'],re.sub(r'\([^)]*\)','',food['name']),food['name'].replace('스프','수프').replace('쥬스','주스')]
    hits=[n+1 for n,s in enumerate(cp) if any(compact(v) in s for v in variants)]
    if hits:r['matches'].append(dict(recipeId=food['recipeId'],name=food['name'],pages=hits))
  except Exception as e:r['error']=str(e)
  return r
 records=[]
 with concurrent.futures.ThreadPoolExecutor(max_workers=4) as pool:
  for r in pool.map(fetch,targets):
   records.append(r);print(json.dumps({k:r.get(k) for k in ('url','status','pageCount','matches','error')},ensure_ascii=False),flush=True)
   (OUT/'source-captures.json').write_text(json.dumps(records,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
if __name__=='__main__':run()
