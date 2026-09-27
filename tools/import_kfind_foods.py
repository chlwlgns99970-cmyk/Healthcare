#!/usr/bin/env python3
"""Convert the official K-FIND food workbook into deterministic Android assets.

The source workbook is never modified. Nutrition values are copied from K-FIND;
missing nutrient values remain empty. Meal-template calories are intentionally not
written as authoritative values because the Android seeder recalculates them from
the referenced FoodItem rows and amounts.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import re
import unicodedata
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

import openpyxl


VERSION = "2026-08-28"
EXPECTED_SOURCE_SHA256 = "1EF3551F9A1D0EE87891D6306FA22BBD6A7FFCC90F2C70A4A184DBFE3FCE6EA6"
EXPECTED_FOOD_COUNT = 19_617
SOURCE_TYPE = "K-FIND"
SOURCE_NAME = "식품영양성분 데이터베이스"
SOURCE_NAME_ENGLISH = "Korean Food Composition Database system(K-FCDB)"
REFERENCE_PATTERN = re.compile(r"^\s*([0-9]+(?:\.[0-9]+)?)\s*(g|ml)\s*$", re.IGNORECASE)
WEIGHT_PATTERN = REFERENCE_PATTERN
NORMALIZED_NAME_PATTERN = re.compile(r"[^0-9a-z가-힣]")


@dataclass(frozen=True)
class TemplateSpec:
    meal_type: str
    mode: str
    food_code: str
    preparation_minutes: int
    cost_level: str
    allergens: tuple[str, ...] = ()
    vegetarian_excluded: bool = True


# Product-curated template metadata. Food names, serving weights and every
# nutrition value are still read from the official workbook by food code.
TEMPLATE_SPECS = (
    # Breakfast: home cooking, dining out, and ready-to-buy choices.
    TemplateSpec("BREAKFAST", "COOK", "D304-180000000-0001", 25, "LOW", ("참깨",), False),
    TemplateSpec("BREAKFAST", "COOK", "D101-044000000-0001", 25, "LOW", ("밀", "대두"), False),
    TemplateSpec("BREAKFAST", "COOK", "D301-002000000-0001", 20, "LOW", ("대두", "밀"), False),
    TemplateSpec("BREAKFAST", "DINING_OUT", "D103-144280000-0001", 5, "MEDIUM", ("메밀", "밀", "대두")),
    TemplateSpec("BREAKFAST", "DINING_OUT", "D102-096280000-0001", 5, "MEDIUM", ("밀", "달걀", "우유", "대두", "참치")),
    TemplateSpec("BREAKFAST", "DINING_OUT", "D101-007520000-0001", 5, "MEDIUM", ("달걀", "대두")),
    TemplateSpec("BREAKFAST", "CONVENIENCE", "D202-096060000-0001", 5, "MEDIUM", ("밀", "달걀", "우유", "대두", "닭고기")),
    TemplateSpec("BREAKFAST", "CONVENIENCE", "D202-115000000-0021", 5, "MEDIUM", ("밀", "달걀", "우유", "돼지고기")),
    TemplateSpec("BREAKFAST", "CONVENIENCE", "D202-083000000-0022", 5, "MEDIUM", ("밀", "달걀", "우유", "대두", "닭고기")),

    # Lunch.
    TemplateSpec("LUNCH", "COOK", "D301-017550000-0001", 25, "LOW", ("새우", "오징어", "조개류", "달걀", "대두")),
    TemplateSpec("LUNCH", "COOK", "D101-028000000-0001", 25, "LOW", ("달걀", "우유", "밀", "대두"), False),
    TemplateSpec("LUNCH", "COOK", "D301-017250000-0001", 25, "LOW", ("새우", "오징어", "조개류", "달걀", "대두")),
    TemplateSpec("LUNCH", "DINING_OUT", "D103-173000000-0001", 5, "MEDIUM", ("밀", "대두"), False),
    TemplateSpec("LUNCH", "DINING_OUT", "D305-256210000-0001", 5, "MEDIUM", ("돼지고기", "대두")),
    TemplateSpec("LUNCH", "DINING_OUT", "D303-164310000-0001", 5, "MEDIUM", ("밀", "새우", "오징어", "조개류", "대두")),
    TemplateSpec("LUNCH", "CONVENIENCE", "D306-286000000-0001", 15, "MEDIUM", ("소고기", "낙지", "대두")),
    TemplateSpec("LUNCH", "CONVENIENCE", "D306-281000000-0001", 15, "MEDIUM", ("소고기", "대두")),
    TemplateSpec("LUNCH", "CONVENIENCE", "D201-059200000-0002", 10, "MEDIUM", ("우유", "대두"), False),

    # Dinner.
    TemplateSpec("DINNER", "COOK", "D306-262000000-0001", 30, "LOW", ("고등어", "대두")),
    TemplateSpec("DINNER", "COOK", "D101-017030000-0001", 20, "LOW", ("달걀", "대두"), False),
    TemplateSpec("DINNER", "COOK", "D103-145000000-0001", 25, "LOW", ("달걀", "소고기", "대두")),
    TemplateSpec("DINNER", "DINING_OUT", "D101-051200000-0001", 5, "MEDIUM", ("생선", "대두")),
    TemplateSpec("DINNER", "DINING_OUT", "D101-042200000-0001", 5, "MEDIUM", ("생선", "새우", "달걀", "대두")),
    TemplateSpec("DINNER", "DINING_OUT", "D103-151300000-0001", 5, "MEDIUM", ("밀", "돼지고기", "소고기", "대두")),
    TemplateSpec("DINNER", "CONVENIENCE", "D202-096000000-0183", 5, "MEDIUM", ("밀", "달걀", "우유", "대두", "닭고기")),
    TemplateSpec("DINNER", "CONVENIENCE", "D203-194000000-0004", 10, "MEDIUM", ("밀", "우유", "대두"), False),
    TemplateSpec("DINNER", "CONVENIENCE", "D210-467000000-0008", 10, "MEDIUM", ("밀", "대두", "어묵")),

    # Snacks.
    TemplateSpec("SNACK", "COOK", "D104-197000000-0001", 20, "LOW", (), False),
    TemplateSpec("SNACK", "COOK", "D120-730000000-0001", 5, "LOW", ("대두",), False),
    TemplateSpec("SNACK", "COOK", "D102-053000000-0001", 10, "LOW", (), False),
    TemplateSpec("SNACK", "DINING_OUT", "D102-072000000-0001", 5, "LOW", (), False),
    TemplateSpec("SNACK", "DINING_OUT", "D102-082290000-0001", 5, "MEDIUM", ("밀", "달걀", "우유", "대두"), False),
    TemplateSpec("SNACK", "DINING_OUT", "D319-711030000-0001", 5, "MEDIUM", ("우유",), False),
    TemplateSpec("SNACK", "CONVENIENCE", "D220-748000000-0209", 5, "LOW", ("우유",), False),
    TemplateSpec("SNACK", "CONVENIENCE", "D220-716000000-0249", 5, "LOW", (), False),
    TemplateSpec("SNACK", "CONVENIENCE", "D202-110000000-0545", 5, "MEDIUM", ("밀", "달걀", "우유"), False),
)


FOOD_HEADERS = (
    "id", "sourceType", "sourceFoodCode", "name", "normalizedName", "aliases", "category",
    "referenceAmount", "unit", "energyKcal", "carbohydrateGrams", "proteinGrams", "fatGrams",
    "sodiumMilligrams", "servingDescription", "brand", "barcode", "dataVersion", "createdAt", "updatedAt",
)
TEMPLATE_HEADERS = (
    "id", "name", "supportedMealTypes", "preparationMinutes", "costLevel", "tags", "allergens",
    "excludedDietTypes", "cuisineType", "source", "createdAt", "updatedAt",
)
INGREDIENT_HEADERS = (
    "id", "mealTemplateId", "foodItemId", "amount", "unit", "adjustable", "minimumAmount",
    "maximumAmount", "adjustmentStep",
)


def normalize_name(value: str) -> str:
    return NORMALIZED_NAME_PATTERN.sub("", unicodedata.normalize("NFKC", value).lower())


def number(value: Any) -> float | None:
    if value is None or str(value).strip() == "":
        return None
    try:
        return float(str(value).replace(",", "").strip())
    except ValueError:
        return None


def compact_number(value: float | None) -> str:
    if value is None:
        return ""
    return f"{value:.10f}".rstrip("0").rstrip(".")


def token_field(values: Iterable[str]) -> str:
    normalized = [value.strip() for value in values if value and value.strip()]
    return f"|{'|'.join(dict.fromkeys(normalized))}|" if normalized else ""


def parse_amount(value: Any, pattern: re.Pattern[str]) -> tuple[float, str]:
    match = pattern.fullmatch(str(value or ""))
    if not match:
        raise ValueError(f"Unsupported amount: {value!r}")
    return float(match.group(1)), match.group(2).lower()


def write_csv(path: Path, headers: tuple[str, ...], rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as output:
        writer = csv.DictWriter(output, fieldnames=headers, extrasaction="ignore", lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def load_source(path: Path) -> tuple[list[dict[str, Any]], dict[str, dict[str, Any]]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if len(workbook.sheetnames) != 1:
        raise ValueError(f"Expected one K-FIND sheet, found {workbook.sheetnames}")
    worksheet = workbook.active
    source_rows = worksheet.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(source_rows)]
    required = {
        "식품코드", "식품명", "대표식품명", "식품대분류명", "영양성분함량기준량", "에너지(kcal)",
        "탄수화물(g)", "단백질(g)", "지방(g)", "나트륨(mg)", "식품중량", "업체명",
        "데이터기준일자",
    }
    missing = sorted(required.difference(headers))
    if missing:
        raise ValueError(f"Required K-FIND columns are missing: {missing}")
    index = {header: position for position, header in enumerate(headers)}

    created_at = int(datetime.fromisoformat(f"{VERSION}T00:00:00+09:00").timestamp() * 1000)
    food_rows: list[dict[str, Any]] = []
    source_by_code: dict[str, dict[str, Any]] = {}
    for excel_row in source_rows:
        code = str(excel_row[index["식품코드"]] or "").strip()
        name = str(excel_row[index["식품명"]] or "").strip()
        if not code or not name:
            raise ValueError("K-FIND row has no food code or food name")
        reference_amount, unit = parse_amount(excel_row[index["영양성분함량기준량"]], REFERENCE_PATTERN)
        energy = number(excel_row[index["에너지(kcal)"]])
        if energy is None or energy < 0:
            raise ValueError(f"Invalid energy for {code}: {excel_row[index['에너지(kcal)']]!r}")
        representative = str(excel_row[index["대표식품명"]] or "").strip()
        aliases = {
            normalize_name(representative),
            *(normalize_name(part) for part in name.split("_") if part.strip()),
        }
        normalized_name = normalize_name(name)
        aliases.discard("")
        aliases.discard(normalized_name)
        brand = str(excel_row[index["업체명"]] or "").strip()
        if brand in {"해당없음", "해당 없음", "-"}:
            brand = ""
        raw_version = str(excel_row[index["데이터기준일자"]] or VERSION).strip()
        data_version = raw_version[:10] if re.match(r"^\d{4}-\d{2}-\d{2}", raw_version) else VERSION
        if data_version != VERSION:
            raise ValueError(f"Unexpected data version for {code}: {data_version}")
        item_id = f"kfind-{code.lower()}"
        row = {
            "id": item_id,
            "sourceType": SOURCE_TYPE,
            "sourceFoodCode": code,
            "name": name,
            "normalizedName": normalized_name,
            "aliases": token_field(sorted(aliases)),
            "category": str(excel_row[index["식품대분류명"]] or "").strip(),
            "referenceAmount": compact_number(reference_amount),
            "unit": unit,
            "energyKcal": compact_number(energy),
            "carbohydrateGrams": compact_number(number(excel_row[index["탄수화물(g)"]])),
            "proteinGrams": compact_number(number(excel_row[index["단백질(g)"]])),
            "fatGrams": compact_number(number(excel_row[index["지방(g)"]])),
            "sodiumMilligrams": compact_number(number(excel_row[index["나트륨(mg)"]])),
            "servingDescription": f"{compact_number(reference_amount)}{unit} 기준",
            "brand": brand,
            "barcode": "",  # The official food workbook has no barcode column.
            "dataVersion": data_version,
            "createdAt": created_at,
            "updatedAt": created_at,
        }
        if code in source_by_code:
            raise ValueError(f"Duplicate K-FIND food code: {code}")
        source_by_code[code] = {
            "food": row,
            "foodWeight": excel_row[index["식품중량"]],
        }
        food_rows.append(row)
    return food_rows, source_by_code


def build_templates(source_by_code: dict[str, dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    template_rows: list[dict[str, Any]] = []
    ingredient_rows: list[dict[str, Any]] = []
    sequence_by_group: dict[tuple[str, str], int] = {}
    for spec in TEMPLATE_SPECS:
        source = source_by_code.get(spec.food_code)
        if source is None:
            raise ValueError(f"Template references missing K-FIND food code: {spec.food_code}")
        food = source["food"]
        amount, unit = parse_amount(source["foodWeight"], WEIGHT_PATTERN)
        if unit != food["unit"]:
            raise ValueError(f"Template unit mismatch for {spec.food_code}: {unit} != {food['unit']}")
        group = (spec.meal_type, spec.mode)
        group_sequence = sequence_by_group.get(group, 0) + 1
        sequence_by_group[group] = group_sequence
        template_id = f"kfind-{spec.meal_type.lower()}-{spec.mode.lower()}-{group_sequence}"
        display_name = food["name"].replace("_", " · ")
        template_rows.append({
            "id": template_id,
            "name": display_name,
            "supportedMealTypes": token_field((spec.meal_type,)),
            "preparationMinutes": spec.preparation_minutes,
            "costLevel": spec.cost_level,
            "tags": token_field((spec.mode,)),
            "allergens": token_field(spec.allergens),
            "excludedDietTypes": token_field(("VEGETARIAN",)) if spec.vegetarian_excluded else "",
            "cuisineType": "",
            "source": f"{SOURCE_NAME} (K-FIND) {VERSION}",
            "createdAt": food["createdAt"],
            "updatedAt": food["updatedAt"],
        })
        step = 10.0 if unit == "ml" or amount >= 100 else 5.0
        ingredient_rows.append({
            "id": "",
            "mealTemplateId": template_id,
            "foodItemId": food["id"],
            "amount": compact_number(amount),
            "unit": unit,
            "adjustable": "true",
            "minimumAmount": compact_number(amount * 0.5),
            "maximumAmount": compact_number(amount * 1.5),
            "adjustmentStep": compact_number(step),
        })

    for group, count in sequence_by_group.items():
        if count < 3:
            raise ValueError(f"Each meal/mode group needs at least three templates: {group} has {count}")
    expected_groups = {(meal, mode) for meal in ("BREAKFAST", "LUNCH", "DINNER", "SNACK") for mode in ("COOK", "DINING_OUT", "CONVENIENCE")}
    if set(sequence_by_group) != expected_groups:
        raise ValueError(f"Template group mismatch: {sorted(set(sequence_by_group) ^ expected_groups)}")

    target_by_meal = {"BREAKFAST": 500, "LUNCH": 700, "DINNER": 600, "SNACK": 200}
    for template, ingredient in zip(template_rows, ingredient_rows, strict=True):
        food = next(value["food"] for value in source_by_code.values() if value["food"]["id"] == ingredient["foodItemId"])
        calories = float(food["energyKcal"]) * float(ingredient["amount"]) / float(food["referenceAmount"])
        target = target_by_meal[template["supportedMealTypes"].strip("|")]
        tolerance = max(target * 0.10, 80)
        if abs(calories - target) > tolerance + 1e-9:
            raise ValueError(f"Template {template['id']} is outside the default calorie tolerance: {calories:.1f} vs {target}")
    return template_rows, ingredient_rows


def main() -> None:
    project_root = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--source",
        type=Path,
        default=project_root / "data-source" / "kfind" / f"kfind-food-db-{VERSION}.xlsx",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=project_root / "app" / "src" / "main" / "assets" / "fooddata",
    )
    args = parser.parse_args()
    source_path = args.source.resolve()
    output_dir = args.output.resolve()
    if not source_path.is_file():
        raise FileNotFoundError(source_path)

    source_sha256 = hashlib.sha256(source_path.read_bytes()).hexdigest().upper()
    if source_sha256 != EXPECTED_SOURCE_SHA256:
        raise ValueError(f"Unexpected K-FIND source SHA-256: {source_sha256}")
    food_rows, source_by_code = load_source(source_path)
    if len(food_rows) != EXPECTED_FOOD_COUNT:
        raise ValueError(f"Expected {EXPECTED_FOOD_COUNT} K-FIND rows, found {len(food_rows)}")
    template_rows, ingredient_rows = build_templates(source_by_code)
    write_csv(output_dir / "food_items.csv", FOOD_HEADERS, food_rows)
    write_csv(output_dir / "meal_templates.csv", TEMPLATE_HEADERS, template_rows)
    write_csv(output_dir / "meal_template_ingredients.csv", INGREDIENT_HEADERS, ingredient_rows)

    bundle_material = f"{source_sha256}|{len(food_rows)}|{len(template_rows)}|{len(ingredient_rows)}|v2"
    bundle_id = hashlib.sha256(bundle_material.encode("utf-8")).hexdigest().upper()
    manifest = (
        f"bundleId={bundle_id}\n"
        f"dataVersion={VERSION}\n"
        f"sourceSha256={source_sha256}\n"
        f"foodCount={len(food_rows)}\n"
        f"templateCount={len(template_rows)}\n"
        f"ingredientCount={len(ingredient_rows)}\n"
        f"sourceName={SOURCE_NAME}\n"
        f"sourceNameEnglish={SOURCE_NAME_ENGLISH}\n"
    )
    (output_dir / "food_data_manifest.properties").write_text(manifest, encoding="utf-8", newline="\n")
    print(
        f"Imported {len(food_rows)} foods, {len(template_rows)} templates and "
        f"{len(ingredient_rows)} ingredients from {source_path.name}."
    )
    print(f"Source SHA-256: {source_sha256}")
    print(f"Bundle ID: {bundle_id}")


if __name__ == "__main__":
    main()
